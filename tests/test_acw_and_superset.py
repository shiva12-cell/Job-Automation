import pytest
from datetime import datetime
import openpyxl
from src.excel_manager import ExcelManager, get_status_style
from src.models import JobApplication, ApplicationStatus
from src.applier.question_superset import QuestionSupersetResolver, QUESTION_SUPERSET_SCHEMA
from src.config import load_candidate_profile


def test_acw_sheet_creation_and_highlighting(tmp_path):
    test_excel = tmp_path / "test_tracker_acw.xlsx"
    mgr = ExcelManager(file_path=str(test_excel))

    today_str = datetime.now().strftime("%Y-%m-%d")
    expected_sheet_name = f"ACW({today_str})"

    app = JobApplication(
        app_id="ACW-TST-001",
        date_applied=today_str,
        company="Microsoft",
        job_title="Data Analyst",
        location="Bengaluru",
        job_url="https://careers.microsoft.com/job/1",
        platform="Company Website",
        status=ApplicationStatus.APPLIED,
        notes="Automated ACW Application",
    )

    created_sheet = mgr.append_acw_application(app, match_pct=88.0, chance="High")
    assert created_sheet == expected_sheet_name

    # Verify workbook contains ACW(Date) sheet
    wb = openpyxl.load_workbook(str(test_excel))
    assert expected_sheet_name in wb.sheetnames

    ws = wb[expected_sheet_name]
    assert ws.max_row >= 2
    # Status column is 8 in standard headers
    status_cell = ws.cell(row=2, column=8)
    assert status_cell.value == "Applied"
    # Verify fill & font highlighting
    assert status_cell.fill.start_color.rgb in ["00C6EFCE", "C6EFCE"]
    assert status_cell.font.bold is True

    # Also verify get_all_rows with "ACW" retrieves the rows
    acw_rows = mgr.get_all_rows(sheet_name="ACW")
    assert len(acw_rows) == 1
    assert acw_rows[0]["Company"] == "Microsoft"
    wb.close()


def test_update_application_status_highlighted(tmp_path):
    test_excel = tmp_path / "test_tracker_update.xlsx"
    mgr = ExcelManager(file_path=str(test_excel))

    app = JobApplication(
        app_id="AUTO-GOOG-999",
        company="Google",
        job_title="Business Analyst",
        platform="Naukri.com",
        status=ApplicationStatus.APPLIED,
    )
    mgr.append_application(app, match_pct=92.0, chance="High")

    # Update status to Interview Scheduled
    updated = mgr.update_application_status_highlighted(
        app_id="AUTO-GOOG-999",
        new_status="Interview Scheduled",
        notes="Round 1 Technical Interview scheduled via recruiter"
    )
    assert updated is True

    # Check that in All Applications and Naukri Jobs, the status cell has interview styling
    wb = openpyxl.load_workbook(str(test_excel))
    for s_name in ["All Applications", "Naukri Jobs", "Business Analyst"]:
        ws = wb[s_name]
        found = False
        for row in range(2, ws.max_row + 1):
            if ws.cell(row=row, column=1).value == "AUTO-GOOG-999":
                status_col = 9 if s_name == "Business Analyst" else 8
                st_cell = ws.cell(row=row, column=status_col)
                assert st_cell.value == "Interview Scheduled"
                assert st_cell.fill.start_color.rgb in ["00D9E1F2", "D9E1F2"]
                assert st_cell.font.color.rgb in ["001F4E79", "1F4E79"]
                found = True
                break
        assert found is True
    wb.close()


def test_question_superset_resolver():
    profile = load_candidate_profile()
    resolver = QuestionSupersetResolver(profile)

    # 1. Test Work Auth
    ans_auth = resolver.resolve_answer("Are you authorized to work in India?")
    assert ans_auth == "Yes"

    ans_spon = resolver.resolve_answer("Will you require visa sponsorship in the future?")
    assert ans_spon == "No"

    # 2. Test Notice Period
    ans_notice = resolver.resolve_answer("What is your official notice period?")
    assert "0" in str(ans_notice) or "immediate" in str(ans_notice).lower()

    # 3. Test Select options
    opt = resolver.resolve_answer(
        "Do you have a disability?",
        options=["Yes, I have a disability", "No, I do not have a disability", "Prefer not to answer"],
        field_type="select"
    )
    assert "No" in opt

    # 4. Schema verification: 12 categories
    assert len(QUESTION_SUPERSET_SCHEMA["categories"]) == 12


def test_question_solver_personal_info():
    from src.applier.question_solver import QuestionSolver
    profile = load_candidate_profile()
    solver = QuestionSolver(profile)

    # Verify candidate identity and contact resolution
    first_name = solver.answer_text_question("First Name")
    assert first_name and first_name != "Yes"
    assert first_name.lower() in profile.personal_info.full_name.lower()

    email = solver.answer_text_question("Email Address")
    assert "@" in email

    phone = solver.answer_text_question("Mobile Number")
    assert phone and phone != "Yes"


def test_redirect_resolver_direct_href():
    from unittest.mock import MagicMock
    from src.applier.naukri_redirect_resolver import resolve_company_site_redirect

    mock_page = MagicMock()
    mock_btn = MagicMock()
    mock_btn.get_attribute.return_value = "https://jobs.lever.co/targetcompany/123"
    mock_page.query_selector.return_value = mock_btn

    mock_context = MagicMock()
    resolved = resolve_company_site_redirect(mock_page, mock_context)
    assert resolved == "https://jobs.lever.co/targetcompany/123"

